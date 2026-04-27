from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import docx


QUESTION_NUMBER_PREFIX = re.compile(r"^\s*\d+\s*[、.．)]\s*")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def strip_question_prefix(text: str) -> str:
    return QUESTION_NUMBER_PREFIX.sub("", text).strip()


def is_question(text: str) -> bool:
    stripped = strip_question_prefix(text)
    if not stripped:
        return False
    return stripped.endswith(("？", "?"))


def extract_blocks(path: Path) -> list[str]:
    document = docx.Document(str(path))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                blocks.append(" | ".join(cells))

    return blocks


def build_pairs(blocks: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    current_question = ""
    answer_parts: list[str] = []

    def flush() -> None:
        nonlocal current_question, answer_parts
        question = strip_question_prefix(current_question)
        answer = "\n".join(part for part in answer_parts if part).strip()
        if question and answer:
            rows.append({"question": question, "answer": answer})
        current_question = ""
        answer_parts = []

    for block in blocks:
        if is_question(block):
            flush()
            current_question = block
            continue

        if current_question:
            answer_parts.append(block)

    flush()
    return rows


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["question", "answer"])
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert FAQ-style DOCX into question/answer CSV.")
    parser.add_argument("--input", required=True, help="Path to the source DOCX file.")
    parser.add_argument("--output", required=True, help="Path to the output CSV file.")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    output_path = Path(args.output).expanduser()

    if not input_path.exists():
        raise FileNotFoundError(f"Input DOCX not found: {input_path}")

    blocks = extract_blocks(input_path)
    rows = build_pairs(blocks)
    if not rows:
        raise RuntimeError("No FAQ pairs were extracted from the DOCX file.")

    write_csv(output_path, rows)
    print(f"Source: {input_path}")
    print(f"Output: {output_path}")
    print(f"FAQ rows: {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
