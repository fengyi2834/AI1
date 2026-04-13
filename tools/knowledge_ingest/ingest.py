import argparse
import csv
import json
import re
import sys
from datetime import datetime, date
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import openpyxl
import docx


SUPPORTED_EXTS = {".xlsx", ".docx"}


def _to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def clean_text(text: str, normalize_spaces: bool) -> str:
    if text is None:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in text.split("\n")]
    cleaned_lines: List[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                cleaned_lines.append("")
            blank = True
            continue
        blank = False
        if normalize_spaces:
            line = re.sub(r"\s+", " ", line)
        cleaned_lines.append(line)
    return "\n".join(cleaned_lines).strip()


def chunk_text(text: str, max_chars: int) -> List[str]:
    if max_chars <= 0 or len(text) <= max_chars:
        return [text]
    lines = text.split("\n")
    chunks: List[str] = []
    current: List[str] = []
    current_len = 0
    for line in lines:
        added = len(line) + (1 if current else 0)
        if current and current_len + added > max_chars:
            chunks.append("\n".join(current))
            current = [line]
            current_len = len(line)
        else:
            current.append(line)
            current_len += added
    if current:
        chunks.append("\n".join(current))
    return chunks


def find_input_files(input_path: Path, recursive: bool) -> List[Path]:
    if input_path.is_file():
        return [input_path]
    if not input_path.is_dir():
        raise FileNotFoundError(f"Input path not found: {input_path}")
    pattern = "**/*" if recursive else "*"
    return [p for p in input_path.glob(pattern) if p.suffix.lower() in SUPPORTED_EXTS]


def extract_docx(file_path: Path, include_tables: bool) -> List[Dict]:
    document = docx.Document(str(file_path))
    records: List[Dict] = []
    index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text or ""
        records.append(
            {
                "block_type": "paragraph",
                "section": "",
                "index": index,
                "content": text,
            }
        )
        index += 1

    if include_tables:
        for table in document.tables:
            for row in table.rows:
                cells = [_to_str(cell.text) for cell in row.cells]
                text = " | ".join(cells)
                records.append(
                    {
                        "block_type": "table_row",
                        "section": "",
                        "index": index,
                        "content": text,
                    }
                )
                index += 1
    return records


def extract_xlsx(file_path: Path) -> List[Dict]:
    workbook = openpyxl.load_workbook(str(file_path), data_only=True)
    records: List[Dict] = []
    index = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [_to_str(cell) for cell in row]
            if all(not value for value in values):
                continue
            text = " | ".join(values)
            records.append(
                {
                    "block_type": "row",
                    "section": sheet.title,
                    "index": index,
                    "content": text,
                }
            )
            index += 1
    return records


def build_records(
    file_path: Path,
    input_root: Path,
    include_tables: bool,
    normalize_spaces: bool,
    max_chars: int,
) -> List[Dict]:
    ext = file_path.suffix.lower()
    if ext == ".docx":
        raw_records = extract_docx(file_path, include_tables)
        source_type = "docx"
    elif ext == ".xlsx":
        raw_records = extract_xlsx(file_path)
        source_type = "xlsx"
    else:
        return []

    relative_source = str(file_path.relative_to(input_root)) if file_path.is_relative_to(input_root) else str(file_path)
    processed: List[Dict] = []
    for record in raw_records:
        cleaned = clean_text(record.get("content", ""), normalize_spaces)
        if not cleaned:
            continue
        for chunk_index, chunk in enumerate(chunk_text(cleaned, max_chars)):
            processed.append(
                {
                    "source_file": relative_source,
                    "source_type": source_type,
                    "block_type": record.get("block_type", ""),
                    "section": record.get("section", ""),
                    "index": record.get("index", 0),
                    "chunk": chunk_index,
                    "content": chunk,
                }
            )
    return processed


def write_markdown(output_path: Path, records: List[Dict]) -> None:
    lines: List[str] = []
    current_section: Optional[str] = None
    for record in records:
        section = record.get("section", "")
        if section and section != current_section:
            if lines:
                lines.append("")
            lines.append(f"## Sheet: {section}")
            current_section = section
        lines.append(record["content"])
        lines.append("")
    output_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def write_json(output_path: Path, records: List[Dict]) -> None:
    output_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(output_path: Path, records: List[Dict]) -> None:
    if not records:
        output_path.write_text("", encoding="utf-8")
        return
    fieldnames = ["source_file", "source_type", "block_type", "section", "index", "chunk", "content"]
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            writer.writerow({name: record.get(name, "") for name in fieldnames})


def parse_formats(value: str) -> List[str]:
    if not value:
        return ["md", "json", "csv"]
    value = value.strip().lower()
    if value == "all":
        return ["md", "json", "csv"]
    return [item.strip() for item in value.split(",") if item.strip()]


def load_config(config_path: Optional[Path]) -> Dict:
    if not config_path:
        return {}
    if not config_path.exists():
        raise FileNotFoundError(f"Config not found: {config_path}")
    return json.loads(config_path.read_text(encoding="utf-8"))


def merge_config(defaults: Dict, config: Dict, cli: Dict) -> Dict:
    merged = dict(defaults)
    merged.update(config)
    for key, value in cli.items():
        if value is not None:
            merged[key] = value
    return merged


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract text from xlsx/docx and export Markdown/JSON/CSV.")
    parser.add_argument("-i", "--input", required=False, help="Input file or directory containing xlsx/docx.")
    parser.add_argument("-o", "--output", required=False, help="Output directory.")
    parser.add_argument("-f", "--format", required=False, help="Output formats: md,json,csv or all.")
    parser.add_argument("--config", required=False, help="Path to JSON config file.")
    parser.add_argument("--recursive", dest="recursive", action="store_true", help="Scan directories recursively.")
    parser.add_argument("--no-recursive", dest="recursive", action="store_false", help="Do not scan recursively.")
    parser.add_argument("--docx-tables", dest="docx_tables", action="store_true", help="Include docx tables.")
    parser.add_argument("--no-docx-tables", dest="docx_tables", action="store_false", help="Skip docx tables.")
    parser.add_argument("--normalize-spaces", dest="normalize_spaces", action="store_true", help="Normalize whitespace.")
    parser.add_argument("--no-normalize-spaces", dest="normalize_spaces", action="store_false", help="Keep whitespace.")
    parser.add_argument("--max-chars", type=int, help="Split content into chunks by max chars. 0 disables.")
    parser.set_defaults(recursive=True, docx_tables=True, normalize_spaces=True)

    args = parser.parse_args()

    defaults = {
        "input": None,
        "output": "out",
        "format": "all",
        "recursive": True,
        "docx_tables": True,
        "normalize_spaces": True,
        "max_chars": 0,
    }

    config = load_config(Path(args.config)) if args.config else {}
    cli_values = {
        "input": args.input,
        "output": args.output,
        "format": args.format,
        "recursive": args.recursive,
        "docx_tables": args.docx_tables,
        "normalize_spaces": args.normalize_spaces,
        "max_chars": args.max_chars,
    }
    settings = merge_config(defaults, config, cli_values)

    if not settings.get("input"):
        parser.print_help()
        return 2

    input_path = Path(settings["input"]).expanduser()
    output_dir = Path(settings["output"]).expanduser()
    formats = parse_formats(settings.get("format", "all"))

    try:
        files = find_input_files(input_path, settings["recursive"])
    except FileNotFoundError as exc:
        print(str(exc), file=sys.stderr)
        return 2

    if not files:
        print("No input files found.", file=sys.stderr)
        return 1

    input_root = input_path if input_path.is_dir() else input_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)

    processed_files = 0
    processed_records = 0
    for file_path in files:
        records = build_records(
            file_path=file_path,
            input_root=input_root,
            include_tables=settings["docx_tables"],
            normalize_spaces=settings["normalize_spaces"],
            max_chars=settings["max_chars"],
        )
        if not records:
            continue
        rel = file_path.relative_to(input_root) if file_path.is_relative_to(input_root) else file_path.name
        base = output_dir / rel
        base.parent.mkdir(parents=True, exist_ok=True)

        if "md" in formats:
            write_markdown(base.with_suffix(".md"), records)
        if "json" in formats:
            write_json(base.with_suffix(".json"), records)
        if "csv" in formats:
            write_csv(base.with_suffix(".csv"), records)

        processed_files += 1
        processed_records += len(records)

    print(f"Processed files: {processed_files}")
    print(f"Output records: {processed_records}")
    print(f"Output dir: {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
